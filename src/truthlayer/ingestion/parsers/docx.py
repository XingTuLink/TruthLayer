"""DOCX parser — python-docx.

Paragraphs are emitted in order with their paragraph index; table rows are
appended afterwards (SOP/policy docs often carry facts in tables).
"""

from __future__ import annotations

from pathlib import Path

from truthlayer.domain.errors import ParserError
from truthlayer.providers.parser_protocol import ParsedBlock


class DocxParser:
    extensions = frozenset({".docx"})

    def parse(self, path: Path) -> list[ParsedBlock]:
        try:
            import docx
        except ImportError as exc:
            raise ParserError(
                "docx support requires the parsers extra: "
                'pip install "truthlayer[parsers]"'
            ) from exc

        try:
            document = docx.Document(path)
        except Exception as exc:
            raise ParserError(f"failed to open docx {path}: {exc}") from exc

        blocks: list[ParsedBlock] = []
        index = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                blocks.append(ParsedBlock(text=text, paragraph=index))
                index += 1

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(ParsedBlock(text=" | ".join(cells), paragraph=index))
                    index += 1
        return blocks
